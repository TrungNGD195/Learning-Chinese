import json
from uuid import uuid4

import pymupdf as fitz
from docx import Document as WordDocument
import pytest

from app.ingestion.parsing.base import ParsingFailure
from app.ingestion.parsing.docx import DocxParser
from app.ingestion.parsing.lesson import LessonSnapshotParser
from app.ingestion.parsing.markdown import MarkdownParser
from app.ingestion.parsing.pdf import PdfParser
from app.ingestion.parsing.text import TxtParser


def test_pdf_is_page_aware_and_reports_no_extractable_text() -> None:
    pdf = fitz.open()
    first = pdf.new_page()
    first.insert_font(fontname="china-s")
    first.insert_text((72, 72), "你好 Chinese", fontname="china-s")
    second = pdf.new_page()
    second.insert_text((72, 72), "Page B text")
    units = PdfParser().parse(pdf.tobytes())
    assert len(units) == 2
    assert units[0].source_locator == {"kind": "PDF_PAGE", "page_number": 1}
    assert "你好" in units[0].text
    assert units[1].source_locator == {"kind": "PDF_PAGE", "page_number": 2}
    assert "Page B text" in units[1].text
    blank = fitz.open()
    blank.new_page()
    with pytest.raises(ParsingFailure, match="No extractable text") as error:
        PdfParser().parse(blank.tobytes())
    assert error.value.code == "NO_EXTRACTABLE_TEXT"


def test_docx_preserves_paragraph_table_paragraph_body_order() -> None:
    doc = WordDocument()
    doc.add_paragraph("before")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    doc.add_paragraph("after")
    from io import BytesIO
    output = BytesIO()
    doc.save(output)
    units = DocxParser().parse(output.getvalue())
    assert [unit.text for unit in units] == ["before", "left\tright", "after"]
    assert units[1].source_locator == {"kind": "DOCX_TABLE_ROW", "block_index": 1, "table_index": 0, "row_index": 0}
    assert DocxParser.version


def test_txt_utf8_bom_and_invalid_bytes() -> None:
    parser = TxtParser()
    units = parser.parse("第一段\n\nSecond".encode("utf-8-sig"))
    assert [item.text for item in units] == ["第一段", "Second"]
    assert units[1].source_locator == {"kind": "TEXT_BLOCK", "block_index": 1}
    with pytest.raises(ParsingFailure) as error:
        parser.parse(b"bad\xff")
    assert error.value.code == "TEXT_DECODING_FAILED"


def test_markdown_source_scanner_preserves_order_and_code_whitespace() -> None:
    raw = "# Heading\n\nParagraph\n\n1. ordered\n- unordered\n\n```py\n  x =  1\n```\n".encode()
    units = MarkdownParser().parse(raw)
    assert [u.source_locator["block_kind"] for u in units] == ["HEADING", "PARAGRAPH", "LIST_ITEM", "LIST_ITEM", "CODE_BLOCK"]
    assert units[0].source_locator["heading_level"] == 1
    assert units[2].source_locator["list_kind"] == "ordered"
    assert units[3].source_locator["list_kind"] == "unordered"
    assert "  x =  1" in units[-1].text
    assert "<h1>" not in " ".join(unit.text for unit in units)


def test_lesson_parser_uses_and_validates_stored_snapshot() -> None:
    lesson_id, hsk_id = str(uuid4()), str(uuid4())
    vocab_a, vocab_b, vocab_link_a, vocab_link_b = [str(uuid4()) for _ in range(4)]
    grammar_a, grammar_b, grammar_link_a, grammar_link_b = [str(uuid4()) for _ in range(4)]
    data = {
        "schema_version": "raw_lesson_snapshot/v1",
        "lesson": {"id": lesson_id, "hsk": {"id": hsk_id, "level_number": 1, "title": "HSK"},
                   "lesson_number": 1, "title": "Stored title", "description": "Stored description", "status": "PUBLISHED"},
        "vocabulary": [
            {"order_index": 0, "lesson_vocabulary_id": vocab_link_a, "vocabulary_id": vocab_a, "hanzi": "甲"},
            {"order_index": 1, "lesson_vocabulary_id": vocab_link_b, "vocabulary_id": vocab_b, "hanzi": "乙"},
        ],
        "grammar": [
            {"order_index": 0, "lesson_grammar_id": grammar_link_a, "grammar_id": grammar_a, "name": "grammar alpha"},
            {"order_index": 1, "lesson_grammar_id": grammar_link_b, "grammar_id": grammar_b, "name": "grammar beta"},
        ],
    }
    units = LessonSnapshotParser().parse(json.dumps(data).encode())
    assert [u.text for u in units] == ["Stored title", "Stored description", "甲", "乙", "grammar alpha", "grammar beta"]
    assert units[2].source_locator["entity_id"] == vocab_a and units[3].source_locator["entity_id"] == vocab_b
    assert units[4].source_locator["entity_id"] == grammar_a and units[5].source_locator["entity_id"] == grammar_b
    assert units[-1].source_locator == {"kind": "LESSON_FIELD", "section": "GRAMMAR", "entity_id": grammar_b,
                                        "field": "name", "order_index": 1}
    data["lesson"]["status"] = "DRAFT"
    with pytest.raises(ParsingFailure) as error:
        LessonSnapshotParser().parse(json.dumps(data).encode())
    assert error.value.code == "INVALID_LESSON_SNAPSHOT"


@pytest.mark.parametrize("raw", [b'{"schema_version":"raw_lesson_snapshot/v1",'])
def test_lesson_parser_rejects_malformed_json(raw: bytes) -> None:
    with pytest.raises(ParsingFailure) as error:
        LessonSnapshotParser().parse(raw)
    assert error.value.code == "INVALID_LESSON_SNAPSHOT"


def test_lesson_parser_rejects_unsupported_snapshot_schema_without_repair() -> None:
    raw = json.dumps({"schema_version": "raw_lesson_snapshot/v2"}).encode()
    with pytest.raises(ParsingFailure) as error:
        LessonSnapshotParser().parse(raw)
    assert error.value.code == "INVALID_LESSON_SNAPSHOT"
